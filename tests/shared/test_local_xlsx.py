"""LocalXlsxAdapter 往返读写测试。"""
from __future__ import annotations

import os

from shared.io_adapter.local_xlsx import LocalXlsxAdapter


def test_write_then_read_roundtrip(tmp_path):
    io = LocalXlsxAdapter()
    target = os.path.join(tmp_path, "out.xlsx")
    rows = [
        {"player_tag": "#A", "name": "甲", "score": 10},
        {"player_tag": "#B", "name": "乙", "score": 20},
    ]
    headers = ["player_tag", "name", "score"]
    io.write_sheet(target, rows, headers=headers)

    read_back = io.read_sheet(target)
    assert len(read_back) == 2
    assert read_back[0]["player_tag"] == "#A"
    assert read_back[1]["name"] == "乙"
    assert read_back[1]["score"] == 20


def test_header_order_preserved(tmp_path):
    io = LocalXlsxAdapter()
    target = os.path.join(tmp_path, "order.xlsx")
    io.write_sheet(target, [{"a": 1, "b": 2, "c": 3}], headers=["c", "b", "a"])
    read_back = io.read_sheet(target)
    assert list(read_back[0].keys()) == ["c", "b", "a"]


def test_empty_rows_skipped(tmp_path):
    io = LocalXlsxAdapter()
    target = os.path.join(tmp_path, "empty.xlsx")
    io.write_sheet(target, [], headers=["x", "y"])
    # 只有表头，无数据行
    assert io.read_sheet(target) == []


def test_fill_merged_propagates_top_left_value(tmp_path):
    # 构造：主信息合并跨行，账号列逐行不同（模拟真实报名表结构）
    from openpyxl import Workbook

    target = os.path.join(tmp_path, "merged.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.append(["提交者", "游戏昵称", "匹配值"])   # 表头
    ws.append(["张三", "账号1", 100])
    ws.append([None, "账号2", 200])              # 归属人合并到上一行
    ws.append([None, "账号3", 300])
    ws.merge_cells("A2:A4")                       # A2 合并覆盖 3 行
    wb.save(target)

    io = LocalXlsxAdapter()
    # 不填充：后两行提交者为空
    plain = io.read_sheet(target, fill_merged=False)
    assert plain[1]["提交者"] in (None, "")
    # 填充：合并值向下填充到每个账号行
    filled = io.read_sheet(target, fill_merged=True)
    assert [r["游戏昵称"] for r in filled] == ["账号1", "账号2", "账号3"]
    assert all(r["提交者"] == "张三" for r in filled)


def test_write_sheet_appends_new_sheet_to_existing_workbook(tmp_path):
    # 先写入源表的一个 sheet，再向同一文件新建"名单"sheet，原 sheet 应保留
    from openpyxl import load_workbook

    io = LocalXlsxAdapter()
    target = os.path.join(tmp_path, "book.xlsx")
    io.write_sheet(target, [{"a": 1}], headers=["a"], sheet="报名")
    io.write_sheet(
        target, [{"rank_order": 1, "player_tag": "#B"}],
        headers=["rank_order", "player_tag"], sheet="名单_2026-07",
    )

    wb = load_workbook(target)
    assert set(wb.sheetnames) == {"报名", "名单_2026-07"}

    # 读回新 sheet 内容
    rows = io.read_sheet(target, sheet="名单_2026-07")
    assert rows[0]["player_tag"] == "#B"

    # 同名 sheet 重复导出：覆盖重建，不新增 sheet
    io.write_sheet(
        target, [{"rank_order": 2, "player_tag": "#C"}],
        headers=["rank_order", "player_tag"], sheet="名单_2026-07",
    )
    wb2 = load_workbook(target)
    assert wb2.sheetnames.count("名单_2026-07") == 1
    rows2 = io.read_sheet(target, sheet="名单_2026-07")
    assert rows2[0]["player_tag"] == "#C"


def test_auto_filter_and_freeze_header(tmp_path):
    from openpyxl import load_workbook

    io = LocalXlsxAdapter()
    target = os.path.join(tmp_path, "view.xlsx")
    io.write_sheet(
        target,
        [{"a": 1, "b": 2}, {"a": 3, "b": 4}],
        headers=["a", "b"],
        sheet="档案",
        auto_filter=True,
        freeze_header=True,
    )
    ws = load_workbook(target)["档案"]
    # 自动筛选覆盖含表头的整个数据区，冻结首行
    assert ws.auto_filter.ref == "A1:B3"
    assert ws.freeze_panes == "A2"


def test_defaults_no_filter_no_freeze(tmp_path):
    from openpyxl import load_workbook

    io = LocalXlsxAdapter()
    target = os.path.join(tmp_path, "plain.xlsx")
    io.write_sheet(target, [{"a": 1}], headers=["a"])
    ws = load_workbook(target).active
    assert ws.auto_filter.ref is None
    assert ws.freeze_panes is None
