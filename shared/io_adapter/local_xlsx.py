"""本地 xlsx 读写实现（基于 openpyxl）。"""
from __future__ import annotations

import os

from openpyxl import Workbook, load_workbook

from shared.io_adapter.base import ExcelIO


class LocalXlsxAdapter(ExcelIO):
    """读写本地 .xlsx 文件。

    读取约定：首行为表头，其余行为数据。空行（全部单元格为空）被跳过。
    """

    def read_sheet(
        self, source: str, sheet: str | None = None, fill_merged: bool = False
    ) -> list[dict]:
        # 需要合并单元格信息时不能用 read_only 模式（该模式下拿不到 merged_cells）
        wb = load_workbook(source, data_only=True, read_only=not fill_merged)
        try:
            ws = wb[sheet] if sheet else wb.active

            if not fill_merged:
                rows_iter = ws.iter_rows(values_only=True)
                try:
                    header_row = next(rows_iter)
                except StopIteration:
                    return []
                headers = [
                    (str(h).strip() if h is not None else f"col{i}")
                    for i, h in enumerate(header_row)
                ]
                result: list[dict] = []
                for raw in rows_iter:
                    if raw is None or all(c is None or str(c).strip() == "" for c in raw):
                        continue
                    row = {}
                    for i, header in enumerate(headers):
                        row[header] = raw[i] if i < len(raw) else None
                    result.append(row)
                return result

            # fill_merged 分支：按合并区把左上值填充到区内所有单元格
            max_row, max_col = ws.max_row, ws.max_column
            if not max_row or not max_col:
                return []
            fill: dict[tuple[int, int], object] = {}
            for mr in ws.merged_cells.ranges:
                top_left = ws.cell(row=mr.min_row, column=mr.min_col).value
                for r in range(mr.min_row, mr.max_row + 1):
                    for c in range(mr.min_col, mr.max_col + 1):
                        fill[(r, c)] = top_left

            headers = [
                (str(ws.cell(row=1, column=c).value).strip()
                 if ws.cell(row=1, column=c).value is not None else f"col{c - 1}")
                for c in range(1, max_col + 1)
            ]
            result = []
            for r in range(2, max_row + 1):
                raw = [ws.cell(row=r, column=c).value for c in range(1, max_col + 1)]
                # 空行判断用原始值，避免被填充值误判为非空
                if all(v is None or str(v).strip() == "" for v in raw):
                    continue
                row = {}
                for c in range(1, max_col + 1):
                    v = raw[c - 1]
                    if (v is None or str(v).strip() == "") and (r, c) in fill:
                        v = fill[(r, c)]
                    row[headers[c - 1]] = v
                result.append(row)
            return result
        finally:
            wb.close()

    def write_sheet(
        self,
        target: str,
        rows: list[dict],
        headers: list[str] | None = None,
        sheet: str | None = None,
        auto_filter: bool = False,
        freeze_header: bool = False,
        highlight_rows: set[int] | None = None,
    ) -> None:
        """写出行字典列表到 xlsx。

        - 若 target 文件已存在：在其中**新建一个 sheet**（保留其余 sheet），
          同名 sheet 会被先删除再重建（便于同月重复导出覆盖旧结果）。
        - 若 target 不存在：新建工作簿，把首个（唯一）sheet 命名为 sheet。
        - auto_filter/freeze_header：见 ExcelIO.write_sheet，用于让阅读者在表格里
          直接按列筛选/排序、并冻结表头。
        """
        if headers is None:
            headers = list(rows[0].keys()) if rows else []

        if os.path.exists(target):
            wb = load_workbook(target)
            title = sheet or "名单"
            if title in wb.sheetnames:
                del wb[title]
            ws = wb.create_sheet(title=title)
        else:
            wb = Workbook()
            ws = wb.active
            if sheet:
                ws.title = sheet

        ws.append(headers)
        for row in rows:
            ws.append([row.get(h) for h in headers])

        # 展示层增强：仅在有表头列时生效
        if headers:
            if auto_filter:
                ws.auto_filter.ref = ws.dimensions
            if freeze_header:
                ws.freeze_panes = "A2"

        wb.save(target)
